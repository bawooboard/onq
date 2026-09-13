"""Knowledge Base — 문서 로딩/청킹/인덱싱/검색.

기본 리트리버: char n-gram TF-IDF (외부 모델 없이 완전 오프라인).
dense(ONQ_RETRIEVER=dense) 선택 시 sentence-transformers + numpy 코사인 유사도.
"""
import hashlib
import json
import pathlib
import re

import numpy as np

try:
    import joblib
except ImportError:  # pragma: no cover
    joblib = None

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

SUPPORTED = {".md", ".txt", ".markdown", ".pdf", ".docx"}


class KnowledgeBase:
    def __init__(self, doc_dir, kb_dir, retriever="tfidf",
                 chunk_size=800, dense_model="intfloat/multilingual-e5-small"):
        self.doc_dir = pathlib.Path(doc_dir)
        self.kb_dir = pathlib.Path(kb_dir)
        self.retriever = retriever
        self.chunk_size = chunk_size
        self.dense_model = dense_model
        self.chunks = []                 # [{"id","source","text"}]
        self._vectorizer = None
        self._mat = None
        self._model = None               # dense 모델 캐시
        self._emb = None                 # dense 임베딩 캐시

    # ── 문서 로딩 ──────────────────────────────────────────────
    def load_file(self, path):
        path = pathlib.Path(path)
        suffix = path.suffix.lower()
        if suffix in (".md", ".txt", ".markdown"):
            for enc in ("utf-8", "cp949"):
                try:
                    return path.read_text(encoding=enc)
                except (UnicodeDecodeError, LookupError):
                    continue
            return path.read_text(errors="replace")
        if suffix == ".pdf":
            from pypdf import PdfReader
            pages = [p.extract_text() or "" for p in PdfReader(str(path)).pages]
            return "\n\n".join(pages)
        if suffix == ".docx":
            import docx
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for tbl in d.tables:
                for row in tbl.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        raise ValueError(f"미지원 형식: {suffix}")

    # ── 청킹 (문단 누적) ───────────────────────────────────────
    def chunk_text(self, text, source):
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        paras = [p.strip() for p in text.split("\n") if p.strip()]
        chunks, cur = [], ""
        for p in paras:
            if cur and len(cur) + len(p) + 1 > self.chunk_size:
                chunks.append(cur)
                cur = p
            else:
                cur = f"{cur}\n{p}" if cur else p
        if cur:
            chunks.append(cur)
        out = []
        for i, c in enumerate(chunks):
            cid = hashlib.md5(f"{source}:{i}".encode()).hexdigest()[:12]
            out.append({"id": cid, "source": source, "text": c})
        return out

    # ── 구축/영속화 ────────────────────────────────────────────
    def count(self):
        return len(self.chunks)

    def doc_names(self):
        return sorted({c["source"] for c in self.chunks})

    def remove_all(self):
        self.chunks = []
        if self.kb_dir.exists():
            for f in ("chunks.json", "index.joblib"):
                p = self.kb_dir / f
                if p.exists():
                    p.unlink()

    def add_text(self, text, source):
        self.chunks.extend(self.chunk_text(text, source))
        self.build_index()
        self.save()

    def rebuild(self, doc_dir=None):
        if doc_dir:
            self.doc_dir = pathlib.Path(doc_dir)
        self.remove_existing_index()
        self.chunks = []
        for p in sorted(self.doc_dir.rglob("*")):
            if p.is_file() and p.suffix.lower() in SUPPORTED:
                try:
                    text = self.load_file(p)
                except Exception as e:
                    print(f"  [skip] {p.name}: {e}")
                    continue
                self.chunks.extend(self.chunk_text(text, p.name))
        self.build_index()
        self.save()
        return {"documents": len(self.doc_names()), "chunks": len(self.chunks)}

    def remove_existing_index(self):
        self._vectorizer, self._mat = None, None

    def build_index(self):
        self._texts = [c["text"] for c in self.chunks]
        self._emb = None
        if not self._texts:
            self._vectorizer, self._mat = None, None
            return
        if self.retriever == "dense":
            self._ensure_dense()
            return
        self._vectorizer = TfidfVectorizer(
            analyzer="char_wb", ngram_range=(2, 4),
            min_df=1, max_df=0.9, sublinear_tf=True)
        self._mat = self._vectorizer.fit_transform(self._texts)

    def _ensure_dense(self):
        if self._model is None:
            from sentence_transformers import SentenceTransformer
            self._model = SentenceTransformer(self.dense_model)
        self._emb = self._model.encode(self._texts, normalize_embeddings=True,
                                       show_progress_bar=False)

    def save(self):
        self.kb_dir.mkdir(parents=True, exist_ok=True)
        (self.kb_dir / "chunks.json").write_text(
            json.dumps(self.chunks, ensure_ascii=False, indent=1), encoding="utf-8")
        if joblib is not None and self._vectorizer is not None:
            joblib.dump({"vec": self._vectorizer, "mat": self._mat},
                        self.kb_dir / "index.joblib")

    def load(self):
        p = self.kb_dir / "chunks.json"
        if p.exists():
            self.chunks = json.loads(p.read_text(encoding="utf-8"))
            self._texts = [c["text"] for c in self.chunks]
            if self.retriever == "dense":
                self._ensure_dense()
            elif (self.kb_dir / "index.joblib").exists() and joblib is not None:
                d = joblib.load(self.kb_dir / "index.joblib")
                self._vectorizer, self._mat = d["vec"], d["mat"]
            else:
                self.build_index()

    # ── 검색 ───────────────────────────────────────────────────
    def search(self, query, k=3, threshold=0.0):
        if not self.chunks:
            return []
        if self._mat is None and self._vectorizer is None:
            self.build_index()
        if self.retriever == "dense" and self._emb is not None:
            qv = self._model.encode([query], normalize_embeddings=True)[0]
            sim = self._emb @ qv
        else:
            qv = self._vectorizer.transform([query])
            sim = cosine_similarity(qv, self._mat)[0]
        idxs = np.argsort(sim)[::-1][:k]
        out = []
        for i in idxs:
            s = float(sim[i])
            if s <= threshold:
                break
            out.append({"source": self.chunks[i]["source"],
                        "score": round(s, 4),
                        "text": self.chunks[i]["text"][:600]})
        return out
